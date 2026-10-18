# hyungchan_hw21.py

#5
import json
python_dict = {
    "name": "홍길동",
    "age": 25,
    "city": "서울"
}
path = r"ch21\json_file\json\hw21_data.json"
with open(path, "w") as file:
    json.dump(python_dict, file, indent=4)
    
print('-------')

#6
import os

# 폴더 지정
folder_path = r"ch21\manage_file"

if os.path.exists(folder_path):
    items = os.listdir(folder_path)
    print(f"{folder_path} 내부의 파일 및 폴더 목록:")
    for item in items:
        print(item)
else:
    print("해당 폴더가 존재하지 않습니다.")


print('-------')

#7
from pathlib import Path

# 폴더 내 특정 확장자 파일 찾기
path = r"./ch21/manage_file"
folder_path = Path(path)
txt_files = list(folder_path.glob("*.txt"))

print("TXT 파일 목록:")
for file in txt_files:
    print(file.name)

print('-------')

#8
from pathlib import Path

# 폴더 생성
new_folder = r"./ch21/test"
folder = Path(new_folder)

if not os.path.exists(new_folder):
    folder.mkdir()
    print(f"{folder} 폴더가 생성되었습니다.")

print('-------')

#9

from openpyxl import Workbook

# 새 엑셀 파일 생성
wb = Workbook()
ws = wb.active
ws.title = "Sheet1"

# 저장
path = r"./ch21/test/data.xlsx"
wb.save(path)
print("엑셀 파일이 생성되었습니다.")

print('-------')

#10

from openpyxl import load_workbook

# 엑셀 파일 로드
path = r"./ch21/test/data.xlsx"
wb = load_workbook(path)
ws = wb.active

# 데이터 추가
ws.append(["이름", "나이", "도시"])
ws.append(["강호동", 55, "서울"])
ws.append(["유재석", 51, "부산"])

# 저장
wb.save(path)
print("데이터가 추가되었습니다.")

print('-------')